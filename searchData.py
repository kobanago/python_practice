from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

path = r"C:\Users\ayumi\Documents\study\python_practice\search.xlsx"
wb = load_workbook(path)
ws = wb["検索"]
wsData = wb["データ"]

# 最終行取得
def getLastRow(ws):
    for row in range(ws.max_row, 1, -1):
        if any(cell.value is not None for cell in ws[row]):
            return row
    return 1

# 最終列取得
def getLastCol(ws):
    for col in range(ws.max_column, 1, -1):
        for cells in ws.iter_cols(min_col=col, max_col=col):
            if any(cell.value is not None for cell in cells):
                return col
    return 1


# 同じ値のセルをリストで取得
def getSameValueRow(value, lastRow, lastCol):
    find_list = []
    for row in ws.iter_rows(1, lastRow, 1, lastCol):
        for cell in row:
            if cell.value == value:
                    find_list.append({"row": cell.row, "col": cell.column})
    return find_list

def exeSearch():
    lastRow = getLastRow(ws)
    lastCol = getLastCol(ws)
    dataLastRow = getLastRow(wsData)

    for i in range(0, dataLastRow, 1):
        targetList = []
        dataValue = wsData.cell(i + 1, 1).value
        targetList = getSameValueRow(dataValue, lastRow, lastCol)
        sameValueNum = len(targetList)
        targetValue = ws.cell(targetList[0]["row"], int(targetList[0]["col"]) + 1).value if targetList else "該当なし"

        wsData.cell(i + 1, 2).value = sameValueNum
        wsData.cell(i + 1, 3).value = targetValue

exeSearch()
wb.save(path)


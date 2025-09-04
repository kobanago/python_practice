from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.common.by import By
# エクセル用
from openpyxl import load_workbook
# 非同期処理用
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait
from webdriver_manager.chrome import ChromeDriverManager

service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service)
wait = WebDriverWait(driver=driver, timeout=20)

# タグ要素取得
def get_target_tag_element(name):
    return driver.find_elements(By.TAG_NAME, name)

# class要素取得
def get_target_class_element(name):
    return driver.find_elements(By.CLASS_NAME, name)

# 要素がクリック可能な状態であるか確認する
def get_target_clickable_element(name):
    el = WebDriverWait(driver, 10).until(
        EC.element_to_be_clickable((By.CLASS_NAME, name))
    )
    if el.is_enabled() != True:
        return None
    elif "disable" in el.get_attribute("class"):
        return None
    else:
        return el

# ランキングから取得
def get_target_from_ranking(data_lists = {
        "rank": [],
        "company_name": [],
        "min_money": [],
        "href": []
    }):
    el = get_target_tag_element("td")
    all_len = len(el)
    href_el = get_target_class_element("TableRowData_nameLink__3gi3h")

    # 2ページ以降の場合
    if len(data_lists["rank"]) != 0:
        data_key = ["rank", "company_name", "min_money", "href"]
        for key in data_key:
            data_lists[key] = data_lists[key]

    for i in range(0, all_len, 6):
        # 順位
        data_lists["rank"].append(el[i].text)

    for j in range(4, all_len, 6):
        # 最低購入金額
        data_lists["min_money"].append(el[j].text)
    
    for el in href_el:
        # 会社名
        data_lists["company_name"].append(el.text)
        # リンク
        data_lists["href"].append(el.get_attribute('href'))
    
    return data_lists

# 必要要素全ページから取得
def get_target_all_page_data():
    all_page_data_list = {}
    data_list = get_target_from_ranking()
    all_page_data_list = data_list

    PAGE_CLASS = "Pagination_next__O7enW"
    
    while True:
        next_page_el = get_target_clickable_element(PAGE_CLASS)
        if next_page_el is None:
            break
        next_page_el.click()
        wait.until(EC.presence_of_all_elements_located)
        data_list = get_target_from_ranking(all_page_data_list)
        all_page_data_list = data_list

    return all_page_data_list

# 優待内容取得
def get_target_from_detail(href_list):
    contents_list = []
    for href in href_list:
        driver.get(href)
        wait.until(EC.presence_of_all_elements_located)
        contents = get_target_tag_element("td")
        contents_list.append(contents[2].text)
    return contents_list

# 必要要素ページ取得
def get_target_all_data():
    data_lists = get_target_all_page_data()
    contents_list = get_target_from_detail(data_lists["href"])
    data_lists["contents"] = contents_list
    return data_lists

# データ書き込み
def input_data(data_dict:dict):
    PATH =  input('input path: ')
    wb = load_workbook(PATH)
    ws = wb["一覧"]

    data_key = ["rank", "company_name", "min_money", "contents"]
    for key_cnt in range(0, len(data_key), 1):
        target_key = data_key[key_cnt]
        target_list = data_dict[target_key]
        list_len = len(target_list)
        for cnt in range(0, list_len, 1):
            ws.cell(cnt + 2, key_cnt + 1).value = target_list[cnt]

    wb.save(PATH)

def exe():
    TARGET_URL = "https://www.nikkei.com/marketdata/ranking-jp/shareholder-benefit/"
    driver.get(TARGET_URL)
    wait.until(EC.presence_of_all_elements_located)

    # 要素取得
    data_list = get_target_all_data()

    # 書き込み
    input_data(data_list)

    driver.quit()

exe()